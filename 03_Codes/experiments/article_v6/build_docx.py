"""Build DisasterAware_IEEESystems_v6.docx on the v2 shell (IEEE Systems template)."""
import copy, os, re, json
from docx import Document
from docx.shared import Pt, Inches, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "/home/claude/repo/02_Article/ieee_systems/DisasterAware_IEEESystems_v2.docx"
FIG = "/home/claude/v6/fig"; MEDIA = "/home/claude/v6/media/word/media"
OUT = "/home/claude/v6/out"; os.makedirs(OUT, exist_ok=True)
MAN = open("/home/claude/v6/v6_manuscript.md", encoding="utf-8").read()
REFS = open("/home/claude/v6/refs.txt", encoding="utf-8").read().splitlines()
BIOS = [b.strip() for b in open("/home/claude/v6/bios.txt", encoding="utf-8").read().split("\n\n") if b.strip() and not b.strip().startswith("[^")]
FIGFILES = {"hindcast": f"{MEDIA}/image12.jpeg", "scalability": f"{MEDIA}/image17.png"}
COL_W = 3.45  # in

TABLES = {
 "I": ("Structural gaps of existing wildfire decision support architectures", ["Gap", "Description"], [
    ["G1", "Predefined alternatives in centralized architectures: the system ranks what it was given and cannot propose what the incident needs."],
    ["G2", "Epistemic uncertainty representable only through prior distributions or belief intervals; reasoning exposed only through post-hoc attribution."],
    ["G3", "Simulation and decision reasoning tightly coupled, so neither is validated on its own and the benefit of a decision mechanism cannot be separated from the forecast behind it."]], [0.5, 2.9]),
 "II": ("Evaluation criteria and their measurement", ["Criterion", "Measurement"], [
    ["Predictive fidelity", "Blind hindcasts against historical fires; front and arrival-time agreement"],
    ["Decision quality", "Paired outcome differences against no action and reference doctrines on identical worlds"],
    ["Timeliness", "Decision-cycle wall time; time to first order"],
    ["Robustness", "Outcome under observation outage and noise"],
    ["Uncertainty handling", "Confidence propagation into every decision"],
    ["Adaptability", "Objects admitted; recurrence across independent worlds"],
    ["Explainability", "Trace completeness of every order, edit, and rejection"],
    ["Scalability", "Cost against agent count, horizon, and grid size"]], [1.1, 2.3]),
 "III": ("Hindcast agreement per historical fire (50 replicates)", ["Metric", "Sc1", "Sc2", "Sc3", "Sc4"], [
    ["Front distance (m)", "329", "305", "113", "391"], ["Arrival-time MAE (h)", "1.46", "1.39", "4.04", "2.58"],
    ["Front agreement", "82 %", "90 %", "89 %", "89 %"], ["Time agreement", "90 %", "98 %", "71 %", "96 %"],
    ["Probability of detection", "92 %", "93 %", "97 %", "85 %"], ["Overall agreement", "88 %", "94 %", "86 %", "90 %"]], [1.4, 0.5, 0.5, 0.5, 0.5]),
 "IV": ("Scenario families (72 × 50 cells at 30 m, 30 paired worlds each)", ["Family", "Landscape", "Pool", "Ignitions", "Observation"], [
    ["S1", "mountain forest, 2 settlements", "full", "1, remote", "full"], ["S2", "mountain forest, 2 settlements", "0.6", "2, remote", "full"],
    ["S3", "rolling interface, 12 settlements", "full", "1, near interface", "full"], ["S4", "rolling interface, 12 settlements", "0.6", "2, near interface", "full"],
    ["S5", "as S4", "0.6", "2, near interface", "outage 0.25, noise 0.10"]], [0.5, 1.2, 0.35, 0.65, 0.75]),
}


def table_v():
    rows = json.load(open("/home/claude/bench/out/report/tables/T1_outcome.csv".replace(".csv", ".json"))) if False else None
    import csv
    R = list(csv.DictReader(open("/home/claude/repo/03_Codes/experiments/article_v4/out_llm/tables/outcomes.csv", encoding="utf-8")))
    by = {(r["scenario"], r["arm"]): r for r in R}
    arms = ["T0", "F5Ev", "DA"]
    out = []
    for sc in ["S1", "S2", "S3", "S4", "S5"]:
        out.append([f"{sc} burned area (ha)"] + [f"{float(by[(sc,a)]['burned']):.1f} ± {float(by[(sc,a)]['burned_hw']):.1f}" for a in arms])
    out.append(["Mean burned area (ha)"] + [f"{float(by[('Avg',a)]['burned']):.1f}" for a in arms])
    out.append(["Extinguished within episode"] + [f"{float(by[('Avg',a)]['success_pct']):.0f} %" for a in arms])
    out.append(["Affected population (persons)"] + [f"{float(by[('Avg',a)]['pop']):.0f}" for a in arms])
    out.append(["Mean J_phys"] + [f"{float(by[('Avg',a)]['j_phys']):.3f}" for a in arms])
    out.append(["Mean J_total"] + [(f"({float(by[('Avg',a)]['j_total']):.3f})" if a == "T0" else f"{float(by[('Avg',a)]['j_total']):.3f}") for a in arms])
    out.append(["Mean J_resp"] + [("-" if a == "T0" else f"{float(by[('Avg',a)]['j_resp']):.3f}") for a in arms])
    return ("Outcome per scenario family and configuration at the end of the episode (mean over 30 paired worlds; ± is the 95 percent half-width)",
            ["Measure", "No DSS", "Closed layer", "DisasterAware"], out, [1.3, 0.6, 0.65, 0.9])


TABLES["V"] = table_v()

doc = Document(SRC)
body = doc.element.body
kids = list(body)
# ---- title
title_p = kids[0]
runs = title_p.findall(qn("w:r"))
for r in runs[1:]: title_p.remove(r)
t = runs[0].find(qn("w:t")); t.text = re.search(r"%TITLE (.*)", MAN).group(1); t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
# ---- remove old body from Abstract (index 7) to the sdt (index 133), keep final sectPr
for k in kids[7:-1]: body.remove(k)
sectPr = kids[-1]


def add_par(text="", style="Text", align=None, size=None, bold=None, italic=None, before=None, after=None, keep_next=False):
    p = doc.add_paragraph(style=style)
    if text: r = p.add_run(text); _fmt(r, size, bold, italic)
    if align is not None: p.alignment = align
    pf = p.paragraph_format
    if before is not None: pf.space_before = Pt(before)
    if after is not None: pf.space_after = Pt(after)
    if keep_next: pf.keep_with_next = True
    return p


def _fmt(r, size=None, bold=None, italic=None, font=None):
    if size: r.font.size = Pt(size)
    if bold is not None: r.font.bold = bold
    if italic is not None: r.font.italic = italic
    if font: r.font.name = font


def add_rich(p, text, size=None):
    """italic for *x* segments; bold for **x**."""
    for seg in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not seg: continue
        if seg.startswith("**"): _fmt(p.add_run(seg[2:-2]), size, bold=True)
        elif seg.startswith("*"): _fmt(p.add_run(seg[1:-1]), size, italic=True)
        else: _fmt(p.add_run(seg), size)


def set_cell_borders(cell, top=None, bottom=None):
    tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for side, val in (("top", top), ("bottom", bottom), ("left", "nil"), ("right", "nil")):
        e = OxmlElement(f"w:{side}")
        if val is None or val == "nil": e.set(qn("w:val"), "nil")
        else: e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(val)); e.set(qn("w:color"), "000000")
        b.append(e)
    tcPr.append(b)


def add_table(num):
    title, header, rows, widths = TABLES[num]
    p = add_par(f"TABLE {num}", style="Table Title", keep_next=True, before=6, after=0)
    p2 = add_par(title, style="Table Title", keep_next=True, after=3);
    for r in p2.runs: r.font.small_caps = True
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header)); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.autofit = False
    tblPr = tbl._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = tblPr.find(qn("w:tblW"))
    if tw is None: tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(sum(widths) * 1440)))
    grid = tbl._tbl.tblGrid
    for gc in list(grid): grid.remove(gc)
    for w in widths:
        gc = OxmlElement("w:gridCol"); gc.set(qn("w:w"), str(int(w * 1440))); grid.append(gc)
    for row in tbl.rows:
        for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    for j, h in enumerate(header):
        c = tbl.rows[0].cells[j]; c.text = ""; pr = c.paragraphs[0]; r = pr.add_run(h); r.font.bold = True; r.font.size = Pt(7.5); pr.paragraph_format.space_after = Pt(0)
        set_cell_borders(c, top=8, bottom=6)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = tbl.rows[i + 1].cells[j]; c.text = ""; pr = c.paragraphs[0]; r = pr.add_run(v); r.font.size = Pt(7.5); pr.paragraph_format.space_after = Pt(0)
            if j > 0 and num in ("III", "V"): pr.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_borders(c, bottom=8 if i == len(rows) - 1 else None)
    add_par("", style="Text", after=2)


def add_figure(name, width_in, caption, wide=False):
    path = FIGFILES.get(name, f"{FIG}/{name}.png")
    if wide:
        # continuous section, one column
        s = doc.add_section(WD_SECTION.CONTINUOUS); _cols(s._sectPr, 1)
    p = doc.add_paragraph(style="Normal"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Inches(width_in))
    c = add_par(caption, style="Figure Caption", after=6)
    if wide:
        c.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        s = doc.add_section(WD_SECTION.CONTINUOUS); _cols(s._sectPr, 2)


def _cols(sp, n):
    cols = sp.find(qn("w:cols"))
    if cols is None: cols = OxmlElement("w:cols"); sp.append(cols)
    for ch in list(cols): cols.remove(ch)
    if n == 1: cols.set(qn("w:num"), "1"); cols.set(qn("w:space"), "288"); cols.attrib.pop(qn("w:equalWidth"), None)
    else:
        cols.set(qn("w:num"), "2"); cols.set(qn("w:space"), "288"); cols.set(qn("w:equalWidth"), "1")
    # keep page geometry of the original section
    for tag in ("w:pgSz", "w:pgMar"):
        if sp.find(qn(tag)) is None: sp.append(copy.deepcopy(sectPr.find(qn(tag))))
    t = sp.find(qn("w:type"))
    if t is None: t = OxmlElement("w:type"); sp.insert(0, t)
    t.set(qn("w:val"), "continuous")


import zipfile, sys
from lxml import etree
EQ_MODE = os.environ.get("EQ_MODE", "omml")
_M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
_eqroot = etree.fromstring(zipfile.ZipFile("/home/claude/v6/eqs.docx").read("word/document.xml"))
_OMATH = [om.find(_M + "oMath") for om in _eqroot.findall(".//" + _M + "oMathPara")]
EQ_LINES = {1: [1, 2], 2: [3, 4], 3: [5], 4: [6], 5: [7], 6: [8], 7: [9], 8: [10], 9: [11], 10: [12], 11: [13, 14], 12: [15, 16], 13: [17, 18]}


def add_equation(num, text):
    lines = EQ_LINES[int(num)]
    for k, ln in enumerate(lines):
        last = k == len(lines) - 1
        p = doc.add_paragraph(style="Equation"); pf = p.paragraph_format; pf.space_before = Pt(3 if k == 0 else 0); pf.space_after = Pt(3 if last else 0); pf.first_line_indent = Inches(0)
        ts = pf.tab_stops; ts.add_tab_stop(Inches(COL_W / 2), WD_TAB_ALIGNMENT.CENTER); ts.add_tab_stop(Inches(COL_W), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run("\t"); r.font.size = Pt(9)
        if EQ_MODE == "omml":
            om = copy.deepcopy(_OMATH[ln - 1])
            for rpr in om.iter(qn("w:rPr")):
                for ch in list(rpr):
                    if ch.tag in (qn("w:sz"), qn("w:szCs")): rpr.remove(ch)
                sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "18"); rpr.append(sz)
            p._p.append(om)
        else:
            from PIL import Image
            img = f"/home/claude/v6/fig/eq/line{ln}.png"; w = Image.open(img).size[0] / 600.0
            p.add_run().add_picture(img, width=Inches(w))
        if last:
            r2 = p.add_run(f"\t({num})"); r2.font.size = Pt(9)


# ---- parse manuscript
abstract = re.search(r"%ABSTRACT (.*)", MAN).group(1); index = re.search(r"%INDEX (.*)", MAN).group(1); ack = re.search(r"%ACK (.*)", MAN).group(1)
p = add_par(style="Abstract"); _fmt(p.add_run("Abstract—"), 9, bold=True, italic=True); _fmt(p.add_run(abstract), 9, bold=True)
p = add_par(style="IndexTerms"); _fmt(p.add_run("Index Terms—"), 9, bold=True, italic=True); _fmt(p.add_run(index), 9, bold=True)

lines = MAN.split("\n")
i = 0
fig_order = []
while i < len(lines):
    ln = lines[i].rstrip()
    i += 1
    if not ln or ln.startswith("%TITLE") or ln.startswith("%AUTHORS") or ln.startswith("%ABSTRACT") or ln.startswith("%INDEX") or ln.startswith("%ACK"): continue
    if ln.startswith("# "):
        txt = ln[2:]
        add_par(txt, style="Heading 1"); continue
    if ln.startswith("## "):
        add_par(ln[3:], style="Heading 2"); continue
    if ln.startswith("%TABLE "):
        add_table(ln.split()[1]); continue
    if ln.startswith("%FIG "):
        name, w, cap = ln[5:].split("|", 2); w = float(w)
        add_figure(name, w, cap, wide=w > 4); fig_order.append(name); continue
    if ln.startswith("%EQ "):
        n, txt = ln[4:].split("|", 1); add_equation(n, txt); continue
    p = add_par(style="Text"); add_rich(p, ln)

add_par("Acknowledgment", style="Heading 1")
add_par(ack, style="Text")
add_par("References", style="Heading 1")
for r in REFS:
    p = doc.add_paragraph(style="References"); p.add_run(r)
    p.paragraph_format.left_indent = Inches(0.25); p.paragraph_format.first_line_indent = Inches(-0.25); p.paragraph_format.space_after = Pt(1)
add_par("", style="Text")
for b in BIOS:
    name = b.split(" received")[0]
    p = doc.add_paragraph(style="References"); r = p.add_run(name); r.font.bold = True; p.add_run(b[len(name):]); p.paragraph_format.space_after = Pt(4)

# header: manuscript line
for sec in doc.sections:
    for hp in sec.header.paragraphs:
        if "REPLACE THIS LINE" in hp.text:
            for r in hp.runs: r.text = ""

out = f"{OUT}/DisasterAware_IEEESystems_v6{'' if EQ_MODE == 'omml' else '_preview'}.docx"
doc.save(out); print("saved", out, "figs", fig_order)
