"""Write the v4 article (benefits of DisasterAware) as tracked changes on the
v2 document.

  python experiments/article_v4/write_article_v4.py \
      --v2 ../02_Article/ieee_systems/DisasterAware_IEEESystems_v2.docx \
      --manuscript ../02_Article/ieee_systems/meta/v4_manuscript.md \
      --results experiments/article_v4/out_thesis --out-dir <dir>

The results directory is what results_v4.py wrote (tables/, figures/,
fill.json). With --source thesis the DisasterAware column comes from the
thesis ladder campaign (stand-in proposer); with --source open from the
LLM campaign. The sentence that says so is chosen here.
"""
from __future__ import annotations

import argparse
import copy
import csv
import json
import os
import re
import shutil
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, os.path.join(ROOT, "experiments"))
sys.path.insert(0, os.path.join(ROOT, "experiments", "open_space"))
import docx_track as T                                   # noqa: E402
import write_article_v3 as V3                            # noqa: E402
from docx import Document                                # noqa: E402
from docx.oxml.ns import qn                              # noqa: E402

T.AUTHOR = "Reviewer"
_new_para, _insert_after, _new_table = V3._new_para, V3._insert_after, V3._new_table
_mark_table_deleted, accept_all = V3._mark_table_deleted, V3.accept_all
parse_manuscript, fill_placeholders, _live_text = V3.parse_manuscript, V3.fill_placeholders, V3._live_text

PROPOSER_SENTENCE = {
    "thesis": ("In this campaign the generative stage of DisasterAware was fed by a "
               "deterministic stand-in proposer that writes rules over the existing "
               "vocabulary and one macro package through the same gate chain, so the "
               "margin of the last rung measures the open decision mechanism with the "
               "least capable proposer; the objects the language-model proposer "
               "composes are read from the recorded runs of Section V-E."),
    "open": ("In this campaign the generative stage of DisasterAware was fed by the "
             "language-model proposer of Section IV-F through the same gate chain, "
             "one proposal per call with a single revision, so the last rung measures "
             "the system as it is meant to run."),
}


def _in_textbox(p):
    a = p.getparent()
    while a is not None:
        if a.tag == qn("w:txbxContent"):
            return True
        a = a.getparent()
    return False


def _body_text(p):
    """Live text of a paragraph without the text of any text box it anchors."""
    out = []
    for t in p.iter(qn("w:t")):
        if _in_textbox(t) or t.getparent().getparent().tag == qn("w:del"):
            continue
        out.append(t.text or "")
    return "".join(out)


def _mark_para_deleted_keep_drawings(p):
    """Tracked-delete the text of a paragraph but keep its anchored drawings
    (figure frames) alive by moving them out before the deletion."""
    kept = []
    for r in list(p.iter(qn("w:r"))):
        if r.find(qn("w:drawing")) is not None or r.find(
                "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent") is not None:
            if r.getparent() is p:
                p.remove(r)
                kept.append(r)
    V3._mark_para_deleted(p)
    return kept


def _retext_caption_boxes(root, seq_no, new_tail):
    """Inside every text box whose caption field is SEQ Fig <seq_no>, delete
    the runs after the field result and insert new_tail (tracked)."""
    hits = 0
    for tb in root.iter(qn("w:txbxContent")):
        for p in tb.iter(qn("w:p")):
            txt = "".join(t.text or "" for t in p.iter(qn("w:t")))
            instr = "".join(t.text or "" for t in p.iter(qn("w:instrText")))
            if "SEQ Fig" not in instr:
                continue
            # the field result is the run holding the number; find fldChar end
            runs = [r for r in p if r.tag == qn("w:r")]
            sep_idx = end_idx = None
            for i, r in enumerate(runs):
                fc = r.find(qn("w:fldChar"))
                if fc is not None and fc.get(qn("w:fldCharType")) == "separate":
                    sep_idx = i
                if fc is not None and fc.get(qn("w:fldCharType")) == "end":
                    end_idx = i
            if end_idx is None or sep_idx is None:
                continue
            num = "".join(t.text or "" for r in runs[sep_idx + 1:end_idx] for t in r.iter(qn("w:t")))
            if num.strip() != str(seq_no):
                continue
            tail_runs = runs[end_idx + 1:]
            for r in tail_runs:
                d = T._el("w:del", **{"w:id": T._nid(), "w:author": T.AUTHOR, "w:date": T.DATE})
                p.replace(r, d)
                d.append(r)
                for t in r.findall(qn("w:t")):
                    t.tag = qn("w:delText")
            ins = T.ins_run(new_tail)
            # copy the formatting of the first deleted run
            if tail_runs:
                rpr0 = tail_runs[0].find(qn("w:rPr"))
                if rpr0 is not None:
                    ins.find(qn("w:r")).insert(0, copy.deepcopy(rpr0))
            p.append(ins)
            hits += 1
    return hits


def table_rows(results):
    rows = list(csv.DictReader(open(os.path.join(results, "tables", "outcomes.csv"), encoding="utf-8")))
    by = {(r["scenario"], r["arm"]): r for r in rows}
    arms = ["T0", "F5", "F5Ev", "DA"]
    header = ["Scenario / measure", "T0", "F5", "F5+Ev", "DisasterAware"]
    out = []
    for sc in ["S1", "S2", "S3", "S4", "S5"]:
        out.append([f"{sc} burned (ha)"] + [f"{float(by[(sc, a)]['burned']):.1f} ± {float(by[(sc, a)]['burned_hw']):.1f}" for a in arms])
    out.append(["Mean burned (ha)"] + [f"{float(by[('Avg', a)]['burned']):.1f}" for a in arms])
    out.append(["Extinguished within episode"] + [f"{float(by[('Avg', a)]['success_pct']):.0f} %" for a in arms])
    out.append(["Affected population"] + [f"{float(by[('Avg', a)]['pop']):.0f}" for a in arms])
    out.append(["Mean Jphys"] + [f"{float(by[('Avg', a)]['j_phys']):.3f}" for a in arms])
    out.append(["Mean Jtotal"] + [f"{float(by[('Avg', a)]['j_total']):.3f}" if a != "T0" else f"({float(by[('Avg', a)]['j_total']):.3f})" for a in arms])
    return header, out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--v2", required=True)
    ap.add_argument("--manuscript", required=True)
    ap.add_argument("--results", required=True)
    ap.add_argument("--source", choices=["thesis", "open"], default=None)
    ap.add_argument("--out-dir", required=True)
    ap.add_argument("--name", default="DisasterAware_IEEESystems_v4")
    args = ap.parse_args()

    F = json.load(open(os.path.join(args.results, "fill.json"), encoding="utf-8"))
    xf = os.path.join(args.results, "extra_fill.json")
    if os.path.exists(xf):
        F.update(json.load(open(xf, encoding="utf-8")))
    source = args.source or F.get("SRC", "thesis")
    F["PROPOSER_SENTENCE"] = PROPOSER_SENTENCE[source]
    M = parse_manuscript(args.manuscript)

    doc = Document(args.v2)
    body = doc.element.body
    paras = [p for p in body.iter(qn("w:p")) if not _in_textbox(p)]
    text = [_body_text(p) for p in paras]

    def find(prefix, start=0):
        for i in range(start, len(paras)):
            if text[i].strip().startswith(prefix):
                return i
        raise KeyError(prefix)

    def para_list(block):
        out = []
        for item in M[block]:
            if isinstance(item, tuple):
                out.append(_new_para(item[1], style="Heading2"))
            else:
                t = fill_placeholders(item, F)
                m = re.match(r"^((?:Open decision space through|Sensitivity analysis with|End-to-end measurement)[^.]*\.)", t)
                out.append(_new_para(t, style="Text", bold_lead=m.group(1) if m else None))
        return out

    def replace_para(idx, new_elems):
        kept = _mark_para_deleted_keep_drawings(paras[idx])
        if kept:
            for r in kept:
                new_elems[0].append(r)
        _insert_after(paras[idx], new_elems)

    # ---- abstract and index terms
    i = find("Abstract")
    ab = [p for p in M["ABSTRACT (replaces paragraph #7)"] if not isinstance(p, tuple)]
    replace_para(i, [_new_para(fill_placeholders(ab[0], F))])
    j = find("Index Terms")
    replace_para(j, [_new_para(ab[1])])

    # ---- introduction #31-#39
    i0 = find("DisasterAware answers these gaps")
    i1 = find("This work is a systems contribution")
    new = para_list("I. INTRODUCTION (paragraphs #31-#39 replaced by the following)")
    for k in range(i0, i1 + 1):
        _mark_para_deleted_keep_drawings(paras[k])
    _insert_after(paras[i1], new)

    # ---- II heading and II.B
    h2 = find("II. Background")
    replace_para(h2, [_new_para("II. Related Work and Background", style="Heading1")])
    hb = find("B. Decision Support Foundations")
    hc = find("C. The Input Space")
    for k in range(hb, hc):
        _mark_para_deleted_keep_drawings(paras[k])
    _insert_after(paras[hc - 1], para_list("II. RELATED WORK AND BACKGROUND (heading replaces \"II. Background\"; II.A kept; II.B replaced by B, C, D below; old II.C, II.D, II.E become E, F, G with their text unchanged)"))
    for old, newh in (("C. The Input Space", "E. The Input Space"), ("D. Concept-Driven", "F. Concept-Driven"), ("E. The Decision Space", "G. The Decision Space")):
        k = find(old)
        replace_para(k, [_new_para(newh + text[k][len(old):], style="Heading2")])

    # ---- V.C setup
    k = find("Synthetic scenarios use")
    replace_para(k, para_list("V.C (replaces the \"C. Experimental Setup and Sensitivity Analysis\" paragraph)"))

    # ---- V.D performance paragraph(s), Table V, Fig. 11 caption
    k = find("Four configurations face identical fires")
    blk = [p for p in M["V.D (replaces the \"Performance of the Decision Layer\" paragraph; Table V and Fig. 11 regenerated)"] if not isinstance(p, tuple)]
    fig_caption = next(p for p in blk if p.startswith("Fig. 11."))
    body_paras = [p for p in blk if not p.startswith("Fig. 11.")]
    replace_para(k, [_new_para(fill_placeholders(p, F), style="Text") for p in body_paras])
    # table V: title line and the table itself
    kt = find("Burned area per scenario and mean decision costs")
    replace_para(kt, [_new_para(fill_placeholders("Outcome per scenario family and configuration at the end of the episode (mean over {{NSEEDS}} paired worlds; ± is the 95 percent half-width)", F), style="TableTitle")])
    tbl = None
    for t in body.iter(qn("w:tbl")):
        first = next(t.iter(qn("w:p")), None)
        if first is not None and _live_text(first).strip().startswith("Scenario"):
            tbl = t
            break
    assert tbl is not None, "Table V not found"
    hdr, rows = table_rows(args.results)
    newt = _new_table(doc, hdr, rows, widths=[1400, 800, 900, 950, 950], size=14)
    _mark_table_deleted(tbl)
    tbl.addnext(newt)
    # figure 11: swap the media and retext the caption
    hits = _retext_caption_boxes(doc.element, 11, "." + fill_placeholders(fig_caption[len("Fig. 11."):], F))
    assert hits >= 1, "Fig. 11 caption not found"

    # ---- V.E products paragraph + heading fix (second "D." -> "E.")
    k = find("Adaptation runs only on demand")
    ve_key = "V.E (heading \"E. Open Intervention Space in Operation\"; first paragraph replaced)"
    fig12_caption = next((p for p in M[ve_key] if not isinstance(p, tuple) and p.startswith("Fig. 12.")), None)
    M[ve_key] = [p for p in M[ve_key] if isinstance(p, tuple) or not p.startswith("Fig. 12.")]
    replace_para(k, para_list(ve_key))
    if fig12_caption:
        hits12 = _retext_caption_boxes(doc.element, 12, "." + fill_placeholders(fig12_caption[len("Fig. 12."):], F))
        assert hits12 >= 1, "Fig. 12 caption not found"
    k = find("D. Open Intervention Space in Operation")
    replace_para(k, [_new_para("E. Open Intervention Space in Operation", style="Heading2")])
    k = find("F. Timeliness and Scalability")
    replace_para(k, [_new_para("F. Timeliness and Scalability", style="Heading2")])
    k = find("E. Compliance with the Evaluation Criteria")
    replace_para(k, [_new_para("G. Compliance with the Evaluation Criteria", style="Heading2")])

    # ---- V.G compliance paragraph
    k = find("Predictive fidelity rests on four blind hindcasts")
    replace_para(k, para_list("V.G (replaces the \"Compliance with the Evaluation Criteria\" paragraph)"))

    # ---- VI
    k1 = find("Three lessons transfer beyond wildfire")
    k2 = find("Three limitations bound the claims")
    new = para_list("VI. DISCUSSION: SYSTEMS LESSONS AND LIMITATIONS (both paragraphs replaced)")
    _mark_para_deleted_keep_drawings(paras[k1])
    _mark_para_deleted_keep_drawings(paras[k2])
    _insert_after(paras[k2], new)

    # ---- VII
    k = find("This article presented DisasterAware")
    replace_para(k, para_list("VII. CONCLUSION AND FUTURE WORK (paragraph replaced)"))

    # ---- references
    last_ref = find("[34]")
    refs = [p for p in M["NEW REFERENCES (appended after [34])"] if not isinstance(p, tuple) and p.startswith("[")]
    new_refs = [_new_para(r, like=paras[last_ref]) for r in refs]
    _insert_after(paras[last_ref], new_refs)
    _sp = paras[last_ref].find(qn("w:pPr"))
    _sect = _sp.find(qn("w:sectPr")) if _sp is not None else None
    if _sect is not None:
        _sp.remove(_sect)
        new_refs[-1].find(qn("w:pPr")).append(_sect)

    # ---- cross references in kept text
    for p in paras:
        t = _live_text(p)
        if p.getparent().tag == qn("w:tc"):
            continue
        t = _body_text(p)
        if "Section II-B" in t or "Section II-C" in t or "Section II-D" in t or "Section II-E" in t:
            t2 = (t.replace("Section II-E", "Section II-G").replace("Section II-D", "Section II-F")
                  .replace("Section II-C", "Section II-E"))
            if t2 != t:
                replace_para(paras.index(p), [_new_para(t2, style="Text")])

    # manuscript-id header line
    for sec in doc.sections:
        for hp in sec.header.paragraphs:
            if "REPLACE THIS LINE" in hp.text:
                for r in hp.runs:
                    r.text = ""

    os.makedirs(args.out_dir, exist_ok=True)
    tracked = os.path.join(args.out_dir, args.name + ".docx")
    doc.save(tracked)
    # swap the Fig. 11 media (image part of rId behind the text box)
    _swap_fig11(tracked, os.path.join(args.results, "figures", "fig11_ladder_burned.png"))
    f12 = os.path.join(args.results, "figures", "fig12_gates.png")
    if fig12_caption and os.path.exists(f12):
        _swap_fig(tracked, 12, f12, fit=True)
    clean = Document(tracked)
    accept_all(clean)
    clean_p = os.path.join(args.out_dir, args.name + "_clean.docx")
    clean.save(clean_p)
    missing = sorted({m for m in re.findall(r"\{\{([^}]+)\}\}", "\n".join(_live_text(p) for p in clean.element.body.iter(qn("w:p"))))})
    print("written", tracked, clean_p)
    print("unfilled placeholders:", missing)


def _swap_fig(path, seq_no, png, fit=True):
    """Replace the image part behind the Fig. <seq_no> text box with png.
    With fit, the picture extent keeps the frame width and takes the height
    of the new image (the frame must not grow, so the image is expected to be
    at least as wide, relative to its height, as the frame)."""
    from lxml import etree
    from PIL import Image
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as zin:
        docxml = zin.read("word/document.xml")
        rels = zin.read("word/_rels/document.xml.rels").decode("utf-8")
        root = etree.fromstring(docxml)
        rid = None
        for tb in root.iter(qn("w:txbxContent")):
            instr = "".join(t.text or "" for t in tb.iter(qn("w:instrText")))
            nums = [t.text for t in tb.iter(qn("w:t")) if (t.text or "").strip().isdigit()]
            if "SEQ Fig" in instr and str(seq_no) in nums:
                blip = next(tb.iter(A + "blip"), None)
                if blip is not None:
                    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if fit:
                        w, h = Image.open(png).size
                        for ext in list(tb.iter(WP + "extent")) + list(tb.iter(A + "ext")):
                            cx = int(ext.get("cx"))
                            ext.set("cy", str(int(cx * h / w)))
                    break
        assert rid, f"Fig. {seq_no} image relationship not found"
        target = re.findall(r'Id="%s"[^>]*Target="([^"]+)"' % rid, rels)[0]
        part = "word/" + target
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == part:
                    data = open(png, "rb").read()
                elif item.filename == "word/document.xml" and fit:
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                zout.writestr(item, data)
    shutil.move(tmp, path)


def _swap_fig11(path, png):
    _swap_fig(path, 11, png, fit=False)


if __name__ == "__main__":
    main()
